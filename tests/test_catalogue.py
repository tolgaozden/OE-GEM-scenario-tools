import openpyxl
import pandas as pd
import pytest

from oegem.catalogue import Catalogue, groups, indicators, search, shock_variables


def test_indicators_and_groups(client, release):
    ind = indicators(client, release["Id"])
    assert list(ind.columns[:2]) == ["code", "name"] and "units" in ind.columns
    assert {"GDP", "CPI", "WPO_EQN", "AAA", "CONF_SHOCK"} <= set(ind["code"])
    assert ind["code"].is_unique
    grp = groups(client, release["Id"])
    assert list(grp.columns[:3]) == ["code", "name", "type"]
    assert {"US", "WORLD", "ALPHA", "BETA"} <= set(grp["code"])


def test_search_options():
    df = pd.DataFrame({"code": ["GDP", "GDPPC", "CPI", "AAA"],
                       "name": ["Real GDP", "GDP per head", "Consumer prices", "Fictional gdp index"]})
    assert list(search(df, "^GDP")["code"]) == ["GDP", "GDPPC"]
    assert list(search(df, "gdp")["code"]) == ["GDP", "GDPPC", "AAA"]
    assert list(search(df, "gdp", case=True)["code"]) == ["AAA"]
    assert list(search(df, "GDP", fields="code")["code"]) == ["GDP", "GDPPC"]
    assert list(search(df, "per (head", regex=False)["code"]) == []
    with pytest.raises(ValueError, match="regex=False"):
        search(df, "per (head")
    with pytest.raises(ValueError, match="columns"):
        search(df, "x", fields=("nope",))


def test_shock_variables(client, release):
    shocks = set(shock_variables(indicators(client, release["Id"]))["code"])
    assert {"CONF_SHOCK", "MMM_SHOCK", "NNN_DUM", "PPP_ADD"} <= shocks  # by code, "dummy", "add factor"
    assert "GDP" not in shocks and "RSH" not in shocks


def test_catalogue_cache_and_labels(client, release, tmp_path):
    cat = Catalogue(client, release["Id"], cache_dir=tmp_path)
    assert "GDP" in set(cat.indicators["code"])
    assert len(cat.groups) == 11
    files = sorted(p.name for p in tmp_path.iterdir())
    assert files == [f"{release['Id']}_groups.json", f"{release['Id']}_indicators.json"]
    n_calls = len(client.session.calls)
    again = Catalogue(client, release["Id"], cache_dir=tmp_path)
    pd.testing.assert_frame_equal(again.indicators, cat.indicators, check_dtype=False)
    assert len(client.session.calls) == n_calls  # served from the JSON cache
    assert again.describe("GDP") == "GDP: Real GDP (synthetic)"
    assert again.describe("GDP:US") == "GDP:US: Real GDP (synthetic), United States (synthetic)"
    assert again.describe("&C:ALPHA").startswith("&C:ALPHA: residual of Private consumption")
    assert again.group_name("NOT_A_GROUP") == "NOT_A_GROUP"
    assert list(again.search("price", kind="indicators")["code"])[:1] == ["CPI"]
    again.refresh()
    assert not list(tmp_path.iterdir())


def test_catalogue_without_disk_cache_and_excel(client, release, tmp_path):
    cat = Catalogue(client, release["Id"], cache_dir=None)
    out = cat.to_excel(tmp_path / "catalogue.xlsx")
    book = openpyxl.load_workbook(out)
    assert book.sheetnames == ["Indicators", "Groups", "Shocks"]
    assert book["Indicators"].max_row == len(cat.indicators) + 1
    assert book["Shocks"].max_row == len(cat.shocks) + 1
    assert list(tmp_path.iterdir()) == [out]
