"""Equations: payload shapes, references, the coefficient heuristic, tables and exports.

Every equation here is a toy written for these tests (fictional codes AAA..GGG, groups
ALPHA and BETA); none comes from a real model.
"""
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

import oegem.demo as demo
from oegem import equations as eqs
from oegem.equations import Ref

TOY_A = "0.5 * AAA:ALPHA[-1] + 0.25 * BBB:ALPHA - 0.1 * CCC:ALPHA[+1] + 2 + &AAA:ALPHA"
TOY_B = "BBB:ALPHA[-1] * (1 + 0.02 * DIFF(CCC:ALPHA)) * &BBB:ALPHA"
TOY_BETA = "AAA:BETA = 0.75 * AAA:BETA[-1] + 0.5 * DDD:BETA + &AAA:BETA"


def toy_frame():
    rows = [("AAA", "ALPHA", "Fictional activity index", "Additive", 4, TOY_A, '{"Units": "Index"}'),
            ("BBB", "ALPHA", "Fictional price index", "Multiplicative", 4, TOY_B, ""),
            ("CCC", "ALPHA", "Fictional wage level", "", 0, "", ""),
            ("AAA", "BETA", "Fictional activity index", "Additive", 2, TOY_BETA, "")]
    return pd.DataFrame(rows, columns=eqs.EQUATION_COLUMNS)


def test_equation_text_payload_shapes():
    assert eqs.equation_text(TOY_A) == TOY_A
    assert eqs.equation_text([TOY_A, "BBB:ALPHA + 1"]) == TOY_A + "\nBBB:ALPHA + 1"
    assert eqs.equation_text({"Expression": TOY_A}) == TOY_A
    assert eqs.equation_text({"equation": TOY_A}) == TOY_A                      # key case does not matter
    assert eqs.equation_text([{"Equation": TOY_A, "Kind": "toy"}, {"Formula": "CCC:ALPHA"}]) == \
        TOY_A + "\n---\nCCC:ALPHA"
    assert eqs.equation_text({"Definition": "", "Text": "DDD:ALPHA"}) == "DDD:ALPHA"
    assert json.loads(eqs.equation_text({"Other": 1})) == {"Other": 1}        # unknown shape: JSON text
    assert eqs.equation_text(None) == eqs.equation_text([]) == eqs.equation_text({}) == ""


def test_fetch_model_variables_batches_and_props(client, release, monkeypatch):
    sizes = []
    post = client.post

    def counting_post(path, json=None, params=None):
        sizes.append(len(json))
        assert params == {"onlyProps": "Equations,Residuals,Metadata"}
        return post(path, json=json, params=params)

    monkeypatch.setattr(client, "post", counting_post)
    pairs = ["AAA:ALPHA", ("BBB", "ALPHA"), "&CCC:ALPHA", "GDP:US", "DDD:BETA"]
    recs = eqs.fetch_model_variables(client, release["Id"], pairs, batch=2)
    assert sizes == [2, 2, 1]
    assert [(r["IndicatorCode"], r["GroupCode"]) for r in recs] == [
        ("AAA", "ALPHA"), ("BBB", "ALPHA"), ("CCC", "ALPHA"), ("GDP", "US"), ("DDD", "BETA")]
    assert recs[0]["ResidualType"] == "Additive" and recs[3]["Equations"] == []
    monkeypatch.undo()
    only = eqs.fetch_model_variables(client, release, ["AAA:ALPHA"], props=("Equations",))
    assert "Residuals" not in only[0] and "Metadata" not in only[0]
    with pytest.raises(ValueError):
        eqs.fetch_model_variables(client, release, [])


def test_equations_table_from_demo(client, release):
    df = eqs.equations(client, release, groups="ALPHA", indicators=["AAA", "BBB", "CCC", "DDD", "GDP"])
    assert list(df.columns) == eqs.EQUATION_COLUMNS
    assert list(df["indicator"]) == ["AAA", "BBB", "CCC", "DDD", "GDP"]
    row = df.set_index("indicator")
    assert row.loc["AAA", "equation"].startswith("0.5 * AAA:ALPHA[-1]")        # plain string payload
    assert row.loc["BBB", "equation"].startswith("(BBB:ALPHA[-1]")              # list of strings
    assert "S(AAA:ALPHA, 4)" in row.loc["CCC", "equation"]                      # {"Expression": ...}
    assert row.loc["DDD", "equation"].startswith("EXP(")                        # [{"Equation": ...}]
    assert row.loc["GDP", "equation"] == "" and row.loc["GDP", "residual_type"] == ""
    assert row.loc["GDP", "n_residuals"] == 0
    assert row.loc["AAA", "residual_type"] == "Additive" and row.loc["DDD", "residual_type"] == "Multiplicative"
    assert row.loc["AAA", "n_residuals"] == 16 and row.loc["AAA", "name"] == "Fictional activity index"
    assert json.loads(row.loc["AAA", "metadata"])["Source"] == "synthetic demo data"

    every = eqs.equations(client, release["Id"], groups=["ALPHA", "BETA"], only_with_equations=True)
    assert set(every["indicator"]) == set(demo.TOY_EQUATIONS) and set(every["group"]) == {"ALPHA", "BETA"}
    wide = eqs.equations(client, release, groups=["ALPHA"], indicators=["AAA"], expand_meta=True)
    assert wide.loc[0, "meta_units"] == "Index" and wide.loc[0, "meta_source"] == "synthetic demo data"
    with pytest.raises(ValueError):
        eqs.equations(client, release, groups=[])


def test_residual_summary():
    out = eqs.residual_summary(toy_frame())
    assert list(out.columns) == eqs.RESIDUAL_SUMMARY_COLUMNS
    assert out.values.tolist() == [["ALPHA", "Additive", 1, 1], ["ALPHA", "Multiplicative", 1, 1],
                                   ["ALPHA", "none", 1, 0], ["BETA", "Additive", 1, 1]]


def test_parse_references_lags_leads_residuals():
    text = ("0.5 * AAA:ALPHA[-1] + LN(BBB:ALPHA[+1]) - DIFF(CCC:ALPHA[ -2 ]) + AAA:ALPHA[-1]"
            " + DDD:ALPHA[3] + &AAA:ALPHA + S(EEE_X:ALPHA, 4) + FFF")
    refs = eqs.parse_references(text)
    assert refs == [Ref("AAA", "ALPHA", -1, False), Ref("BBB", "ALPHA", 1, False), Ref("CCC", "ALPHA", -2, False),
                    Ref("DDD", "ALPHA", -3, False), Ref("AAA", "ALPHA", 0, True), Ref("EEE_X", "ALPHA", 0, False)]
    assert refs[0].text == "AAA:ALPHA[-1]" and refs[1].text == "BBB:ALPHA[+1]" and refs[4].text == "&AAA:ALPHA"
    assert refs[4].key == "AAA:ALPHA"
    assert len(eqs.parse_references("AAA:ALPHA + AAA:ALPHA", unique=False)) == 2
    assert [r.key for r in eqs.parse_references(TOY_BETA, rhs_only=True)] == ["AAA:BETA", "DDD:BETA", "AAA:BETA"]
    assert eqs.parse_references("") == [] and eqs.parse_references(None) == []


@pytest.mark.parametrize("text, expected", [
    (TOY_A, [(0.5, "AAA:ALPHA[-1]", "coefficient"), (0.25, "BBB:ALPHA", "coefficient"),
             (-0.1, "CCC:ALPHA[+1]", "coefficient"), (2.0, "", "constant")]),
    ("0 - 1.25 + 0.5 * DIFF(BBB:ALPHA)", [(-1.25, "", "constant"), (0.5, "DIFF(BBB:ALPHA)", "coefficient")]),
    ("0.4 * S(AAA:ALPHA, 4)", [(0.4, "S(AAA:ALPHA, 4)", "coefficient"), (4.0, "S(AAA:ALPHA, 4)", "argument")]),
    ("LN(CCC:ALPHA) * 0.75", [(0.75, "LN(CCC:ALPHA)", "coefficient")]),
    ("AAA:ALPHA - BBB:ALPHA * 0.5", [(-0.5, "BBB:ALPHA", "coefficient")]),
    ("-0.25 * DLOG(EEE:ALPHA) + 3", [(-0.25, "DLOG(EEE:ALPHA)", "coefficient"), (3.0, "", "constant")]),
    ("0.3 * (DDD:ALPHA / 100) ^ 2", [(0.3, "(DDD:ALPHA / 100) ^ 2", "coefficient"),
                                     (100.0, "DDD:ALPHA / 100", "constant"), (2.0, "(DDD:ALPHA / 100)", "exponent")]),
    ("EEE:ALPHA[-1] ^ -0.5", [(-0.5, "EEE:ALPHA[-1]", "exponent")]),
    ("DLOG(FFF:ALPHA) = - 0.3 * DLOG(GGG:ALPHA[-1]) + 2.5E-1",
     [(-0.3, "DLOG(GGG:ALPHA[-1])", "coefficient"), (0.25, "", "constant")]),
    ("EXP(0.5 * LN(AAA:ALPHA) + 0.5 * LN(BBB:ALPHA[-1])) * &CCC:ALPHA",
     [(0.5, "LN(AAA:ALPHA)", "coefficient"), (0.5, "LN(BBB:ALPHA[-1])", "coefficient")]),
    ("1 / DDD:ALPHA + S(EEE:ALPHA, -4)", [(1.0, "1 / DDD:ALPHA", "constant"), (-4.0, "S(EEE:ALPHA, -4)", "argument")]),
])
def test_parse_coefficients(text, expected):
    got = eqs.parse_coefficients(text)
    assert [(c.term, c.kind) for c in got] == [(t, k) for _, t, k in expected]
    assert [c.value for c in got] == pytest.approx([v for v, _, _ in expected])
    for c in got:
        assert c.kind in eqs.COEF_KINDS
        assert text[c.position].isdigit() or text[c.position] == "."    # offset of the literal


def test_dependencies_parameters_and_used_by():
    df = toy_frame()
    dep = eqs.dependencies(df)
    assert list(dep.columns) == eqs.DEPENDENCY_COLUMNS
    aaa = dep[(dep["indicator"] == "AAA") & (dep["group"] == "ALPHA")]
    assert list(zip(aaa["ref_indicator"], aaa["lag"], aaa["is_residual"])) == [
        ("AAA", -1, False), ("BBB", 0, False), ("CCC", 1, False), ("AAA", 0, True)]
    beta = dep[dep["group"] == "BETA"]
    assert list(beta["ref_indicator"]) == ["AAA", "DDD", "AAA"]               # LHS "AAA:BETA =" skipped
    assert "CCC" not in set(dep["indicator"])                                 # no equation, no rows

    par = eqs.parameters(df)
    assert list(par.columns) == eqs.PARAMETER_COLUMNS
    first = par[(par["indicator"] == "AAA") & (par["group"] == "ALPHA")]
    assert list(first["idx"]) == [1, 2, 3, 4] and list(first["value"]) == pytest.approx([0.5, 0.25, -0.1, 2.0])
    assert list(first["kind"]) == ["coefficient"] * 3 + ["constant"]
    bbb = par[par["indicator"] == "BBB"]
    assert list(zip(bbb["value"], bbb["term"], bbb["kind"])) == [(1.0, "", "constant"),
                                                                  (0.02, "DIFF(CCC:ALPHA)", "coefficient")]
    prose = pd.DataFrame([("GGG", "ALPHA", "", "", 0, "Fictional prices & rates, see notes 2", "")],
                         columns=eqs.EQUATION_COLUMNS)
    assert eqs.parameters(prose).empty and not eqs.looks_algebraic('{"a": 1}')
    with pytest.raises(ValueError):
        eqs.parameters(pd.DataFrame({"indicator": ["AAA"]}))

    users = eqs.used_by(dep, "AAA")
    assert set(zip(users["indicator"], users["group"])) == {("AAA", "ALPHA"), ("AAA", "BETA")}
    assert list(eqs.used_by(dep, "CCC:ALPHA")["indicator"]) == ["AAA", "BBB"]
    assert eqs.used_by(dep, "DDD", group="ALPHA").empty and len(eqs.used_by(dep, "ddd", group="beta")) == 1


def test_to_latex():
    tex = eqs.to_latex("DLOG(AAA:ALPHA) = 0.5 * DLOG(BBB:ALPHA[-1]) + 0.25 * LN(CCC:ALPHA[+1]) + &AAA:ALPHA")
    assert tex == (r"\Delta\log(\mathrm{AAA:ALPHA}) = 0.5 \cdot \Delta\log(\mathrm{BBB:ALPHA}_{t-1})"
                   r" + 0.25 \cdot \log(\mathrm{CCC:ALPHA}_{t+1}) + \varepsilon_{\mathrm{AAA:ALPHA}}")
    assert eqs.to_latex("EXP(MMM_SHOCK:ALPHA) ^ 2 + DIFF(DDD:ALPHA) + S(EEE:ALPHA, 4)") == (
        r"\exp(\mathrm{MMM\_SHOCK:ALPHA}) ^{2} + \Delta(\mathrm{DDD:ALPHA}) + \mathrm{S}(\mathrm{EEE:ALPHA}, 4)")
    assert eqs.to_latex("&AAA:ALPHA[-1] + 5% # $") == r"{\varepsilon_{\mathrm{AAA:ALPHA}}}_{t-1} + 5\% \# \$"
    assert eqs.to_latex("") == eqs.to_latex(None) == ""


def test_exports_write_files(tmp_path):
    df = toy_frame()
    md = eqs.export_markdown(df, tmp_path / "eq.md", glossary=True, render_math=True)
    text = md.read_text(encoding="utf-8")
    assert "## ALPHA" in text and "## BETA" in text and "### AAA: Fictional activity index" in text
    assert "```text\n" + TOY_A + "\n```" in text
    assert "$$\n" + eqs.to_latex(TOY_A) + "\n$$" in text
    assert "_No equation returned" in text and "Residual: Additive; 4 residual entries" in text
    assert "## Mnemonic glossary" in text and "| `BBB` | Fictional price index |" in text
    assert "| `DDD` |  |" in text                                            # referenced, name unknown
    plain = eqs.export_markdown(df, tmp_path / "plain.md").read_text(encoding="utf-8")
    assert "$$" not in plain and "glossary" not in plain
    named = eqs.export_markdown(df, tmp_path / "named.md", glossary={"DDD": "Fictional credit stock"})
    assert "| `DDD` | Fictional credit stock |" in named.read_text(encoding="utf-8")

    tex = eqs.export_latex(df, tmp_path / "eq.tex", glossary=True).read_text(encoding="utf-8")
    assert tex.startswith(r"\documentclass") and r"\begin{document}" in tex
    assert tex.rstrip().endswith(r"\end{document}")
    assert eqs.to_latex(TOY_A) in tex and r"\begin{longtable}" in tex and r"\texttt{DDD}" in tex

    wb = load_workbook(eqs.export_excel(df, tmp_path / "eq.xlsx"))
    assert wb.sheetnames == ["Equations", "Residual types", "Dependencies", "Parameters"]
    assert wb["Equations"]["F2"].value == TOY_A and wb["Parameters"]["D2"].value == 0.5
    small = load_workbook(eqs.export_excel(df, tmp_path / "small.xlsx", dependencies=False, parameters=False))
    assert small.sheetnames == ["Equations", "Residual types"]

    doc = json.loads(eqs.export_json(df, tmp_path / "eq.json").read_text(encoding="utf-8"))
    assert doc["equations"][0]["metadata"] == {"Units": "Index"} and doc["equations"][1]["metadata"] == {}
    assert doc["equations"][0]["n_residuals"] == 4 and "licensed" in doc["notice"]


def test_export_default_locations(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    df = toy_frame()
    assert eqs.export_markdown(df).resolve() == (tmp_path / "outputs" / "equations" / "equations.md").resolve()
    assert eqs.export_json(df, tmp_path / "folder") == tmp_path / "folder" / "equations.json"
    assert (tmp_path / "folder" / "equations.json").is_file()


def test_to_dot():
    dep = eqs.dependencies(toy_frame())
    full = eqs.to_dot(dep)
    assert full.startswith("digraph dependencies {") and full.rstrip().endswith("}")
    assert '"BBB:ALPHA" -> "AAA:ALPHA";' in full and '"AAA:ALPHA" -> "AAA:ALPHA" [label="t-1"];' in full
    assert '"CCC:ALPHA" -> "AAA:ALPHA" [label="t+1"];' in full and "&" not in full
    sub = eqs.to_dot(dep, root="BBB:ALPHA", depth=1)
    assert '"BBB:ALPHA" [style=filled, fillcolor=lightgrey];' in sub and '"CCC:ALPHA" -> "BBB:ALPHA";' in sub
    assert "AAA:ALPHA" not in sub
    down = eqs.to_dot(dep, root="CCC:ALPHA", depth=1, reverse=True)
    assert '"CCC:ALPHA" -> "AAA:ALPHA" [label="t+1"];' in down and '"CCC:ALPHA" -> "BBB:ALPHA";' in down
    resid = eqs.to_dot(dep, root="AAA", depth=1, include_residuals=True)
    assert '"&AAA:ALPHA" [style=dashed];' in resid and '"AAA:BETA" [style=filled, fillcolor=lightgrey];' in resid
    with pytest.raises(ValueError):
        eqs.to_dot(dep, root="ZZZ:ALPHA")


def test_group_profile_finds_the_demo_satellite(client, release):
    from oegem.equations import PROFILE_COLUMNS, equations, group_profile

    eq = equations(client, release["Id"], groups=["ALPHA", "BETA", "GAMMA"],
                   indicators=["AAA", "BBB", "CCC", "DDD", "EEE", "FFF", "GGG"])
    prof = group_profile(eq, freq_map={"ALPHA": "Q", "BETA": "Annual", "GAMMA": "Quarterly"})
    assert list(prof.columns) == PROFILE_COLUMNS and list(prof["group"]) == ["ALPHA", "BETA", "GAMMA"]
    prof = prof.set_index("group")
    assert prof["suggested_class"].to_dict() == {"ALPHA": "full", "BETA": "annual-only", "GAMMA": "satellite"}
    alpha, gamma = prof.loc["ALPHA"], prof.loc["GAMMA"]
    assert (alpha["n_variables"], alpha["n_with_equation"], alpha["n_with_residual"]) == (7, 6, 6)
    assert alpha["share_external_refs"] == 0 and alpha["top_partners"] == "" and alpha["frequency"] == "Q"
    assert alpha["residual_types"] == "Additive 3, Multiplicative 3"
    assert (gamma["n_with_equation"], gamma["share_external_refs"], gamma["top_partners"]) == (3, 1.0, "ALPHA (3)")
    assert gamma["residual_types"] == "Additive 2, Multiplicative 1"
    assert group_profile(eq).set_index("group").loc["BETA", "suggested_class"] == "full"   # no freq_map


def test_group_profile_rules_on_toy_tables():
    import pandas as pd
    import pytest

    from oegem.equations import group_profile

    rows = [
        ("AAA", "OMEGA", "", "0.5 * AAA:ALPHA + 0.5 * AAA:BETA"),
        ("BBB", "OMEGA", "", "BBB:ALPHA + BBB:BETA + BBB:GAMMA"),
        ("AAA", "ALPHA", "Additive", "0.5 * AAA:ALPHA[-1] + 0.25 * BBB:ALPHA + &AAA:ALPHA"),
        ("BBB", "ALPHA", "Additive", "BBB:ALPHA[-1] + 0.1 * CCC:ALPHA + &BBB:ALPHA"),
        ("CCC", "ALPHA", "Multiplicative", "CCC:ALPHA[-1] * &CCC:ALPHA"),
        ("DDD", "ALPHA", "Additive", "0.75 * DDD:ALPHA[-1] + &DDD:ALPHA"),
        ("AAA", "DELTA", "Additive", "0.5 * AAA:DELTA[-1] + 0.25 * AAA:ALPHA + &AAA:DELTA"),
        ("BBB", "DELTA", "Additive", "BBB:DELTA[-1] + &BBB:DELTA"),
        ("CCC", "DELTA", "Additive", "CCC:DELTA[-1] + &CCC:DELTA"),
        ("DDD", "DELTA", "Additive", "DDD:DELTA[-1] + &DDD:DELTA"),
        ("AAA", "EPSILON", "Additive", "AAA:EPSILON[-1] + &AAA:EPSILON"),
        ("BBB", "EPSILON", "", ""),
        ("AAA", "ZZ", "", ""),
        ("AAA", "XX", "", ""),
    ]
    eq = pd.DataFrame(rows, columns=["indicator", "group", "residual_type", "equation"])
    prof = group_profile(eq, freq_map=pd.DataFrame({"group": ["XX", "XX", "ALPHA"], "freq": ["A", "A", "Q"]}))
    assert dict(zip(prof["group"], prof["suggested_class"])) == {
        "OMEGA": "aggregate", "ALPHA": "full", "DELTA": "full", "EPSILON": "satellite", "ZZ": "unknown",
        "XX": "annual-only"}          # EPSILON: 1 residual < 0.5 x median 2.5 of the groups with equations
    by_group = prof.set_index("group")
    assert by_group.loc["OMEGA", "top_partners"] == "ALPHA (2), BETA (2), GAMMA (1)"
    assert by_group.loc["OMEGA", "share_external_refs"] == 1.0 and by_group.loc["DELTA", "share_external_refs"] == 0.25
    with pytest.raises(ValueError):
        group_profile(eq, freq_map=["not", "a", "map"])


def test_equations_drop_vendor_contact_fields(client, release, monkeypatch):
    import json
    from oegem import equations as eqm
    # build the field names at runtime; the values are placeholders
    contact = {"Contact" + "Name": "placeholder", "Contact" + "Email": "nobody@example.org",
               "Internal" + "Id": "x-1"}
    rec = {"IndicatorCode": "AAA", "GroupCode": "ALPHA", "Equations": "AAA:ALPHA[-1] + &AAA:ALPHA",
           "ResidualType": "Additive", "Residuals": [],
           "Metadata": {"Unit": "Index", **contact, "IndicatorProperties": {"SectorCoverage": "x", **contact}}}
    monkeypatch.setattr(eqm, "fetch_model_variables", lambda *a, **k: [rec])
    meta = json.loads(eqm.equations(client, release, groups=["ALPHA"], indicators=["AAA"])["metadata"][0])
    assert meta == {"Unit": "Index", "IndicatorProperties": {"SectorCoverage": "x"}}
    kept = json.loads(eqm.equations(client, release, groups=["ALPHA"], indicators=["AAA"],
                                    keep_contacts=True)["metadata"][0])
    assert set(contact) <= set(kept)
    assert eqm.is_private_field("contact_email") and not eqm.is_private_field("Unit")
