"""Scenario report tables, and writers for xlsx, csv, markdown and html.

`build_report` is the engine behind ``oegem report REPORT.yaml``. A report config::

    title: Oil price scenarios          # optional
    input: outputs/example_oil          # a run directory (holding results.*) or a results
                                        # file, relative to this config file
    # ...or, instead of input, any forecasts in your workspace or the release tree:
    # forecasts:
    #   baseline: latest:5yr
    #   oil_up: /me/oegem-runs/example_oil/oil_up-20261015-101500
    # fetch: {indicators: [GDP, CPI], groups: [US, CHINA], from: 2024Q1, to: 2030Q4}
    #   (fetch also takes freq, annual_fallback and root, the release tree root)
    baseline: baseline                  # label of the baseline forecast
    scenarios: all                      # or a list of labels
    groups: [US, CHINA]                 # table rows; default: every group in the data
    aggregates: groupings.yaml          # optional, see oegem.aggregate (or inline {AGG: {GRP: w}})
    decimals: 2
    variables:
      - indicator: GDP
        label: Real GDP growth
        quarterly: [yoy]                # chain for the quarterly tables (omit for none)
        annual: [annual_avg, yoy]       # chain for the annual tables (omit for none)
        deviation: abs                  # abs (absolute / pp difference) | pct (% of baseline) | none
      - {indicator: RSH, label: Short-term interest rate, quarterly: [level], units: "%"}
    periods: {quarterly: "2026Q1:2028Q4", annual: "2026:2030"}   # quote ranges in YAML
    outputs: {dir: report, xlsx: true, csv: false, markdown: true, html: true, charts: true}
                                        # also name: the file stem (default report)
    mark_forecast: false                # true: quarterly column headers after the history end get "f"
    metadata: true                      # false: do not read or fetch variable metadata

GDP, CPI, RSH, US and CHINA are examples: check every mnemonic in your own catalogue
(`oegem catalogue search`).

Variable metadata (oegem.metadata) comes from the run directory's metadata.csv (written
by `oegem run`) or, with `forecasts:`, is fetched from the baseline. With it the README
lists each indicator's unit, scaling and history end by group, the charts draw a thin
line at the history end and shade the forecast period, the transform step `annual`
annualises each series with its own method (sum, mean or end of year), and
``mark_forecast: true`` appends "f" to quarterly column headers later than the earliest
history end among a table's groups. Without metadata the report is built as before.

Per variable the report builds up to four tables, named ``<NAME>_Q``, ``<NAME>_Q_dev``,
``<NAME>_A`` and ``<NAME>_A_dev`` (NAME = ``name:`` or the indicator; at most 31
characters, unique). Rows are (group, forecast), aggregates appended as extra groups;
columns are the periods inside ``periods:``. Other variable keys: ``name`` (table name
prefix), ``units`` (free text; "%", "percent", "rate" or "pp" marks a rate, so absolute
differences are labelled pp), ``decimals``. An annual chain without an annual_* step gets
``annual_avg`` in front. When ``deviation`` is omitted it is ``abs`` for growth chains and
rates and ``pct`` for level chains. A "README" table comes first: source, baseline,
scenarios, every table's measure chain and units notes.

Transforms run on the full history (so lags exist) before the period filter, and
aggregates are weighted means of the transformed member values and deviations, never of
levels.
"""
from __future__ import annotations

import copy
import datetime as _dt
import html as _html
import json
import logging
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
import yaml

from . import transforms
from .aggregate import as_groupings, weighted_mean
from .data import as_long_frame, empty_frame
from .metadata import METADATA_FILE, _end_key, history_end, indicator_notes, load_metadata
from .metadata import metadata as fetch_metadata
from .periods import freq_of, in_range, normalize_period, parse_period, parse_range, period_key

log = logging.getLogger("oegem")

DEVIATION_KINDS = ("abs", "pct")
RESULT_FILES = ("results.csv", "results.parquet", "results.pq", "results.xlsx")
README = "README"
_RATE_UNITS = ("%", "percent", "per cent", "pct", "rate", "pp", "percentage points", "% pa", "% p.a.")
_SHEET_BAD = re.compile(r"[\[\]:*?/\\']")
_FREQ_WORD = {"Q": "quarterly", "A": "annual"}
_NOTES = {
    "level": "level = values as published by the model, in the indicator's own units.",
    "yoy": "yoy = % change on the same period a year earlier (annual series: on the previous year).",
    "qoq": "qoq = % change on the previous quarter.",
    "qoq_ann": "qoq_ann = % change on the previous quarter, compounded to an annual rate.",
    "diff": "diff = change on the previous period (diffN: on N periods earlier), in the series' own units.",
    "annual_avg": "annual_avg = calendar-year mean of quarterly values; years with fewer than 4 quarters are "
                  "blank; annual-only series pass through unchanged.",
    "annual_sum": "annual_sum = calendar-year sum of quarterly values (complete years only).",
    "annual_last": "annual_last = value in the last quarter of the year.",
    "dev_abs": "dev_abs on a rate = percentage-point difference: scenario minus baseline, in the units of the "
               "measure (pp for rates and growth rates, own units for levels).",
    "dev_pct": "dev_pct = % difference from baseline: (scenario / baseline - 1) x 100.",
}
_FORECAST_NOTE = ("A trailing f on a quarterly column header marks a forecast period: later than the earliest "
                  "history end among the table's groups (variable metadata).")


# =========================================================================== config
@dataclass
class ReportVariable:
    """One `variables:` entry of a report config, normalised."""

    indicator: str
    label: str
    name: str
    quarterly: Optional[List[str]] = None
    annual: Optional[List[str]] = None
    deviation: Optional[str] = "auto"      # "abs" | "pct" | None (no deviation tables) | "auto"
    units: str = ""
    rate: bool = False
    decimals: Optional[int] = None

    def chain(self, freq: str) -> Optional[List[str]]:
        """Transform chain for "Q" or "A" (None when that frequency is not reported)."""
        return self.quarterly if freq == "Q" else self.annual

    def deviation_kind(self, freq: str) -> Optional[str]:
        """"abs", "pct" or None for the chain of this frequency."""
        if self.deviation != "auto":
            return self.deviation
        chain = self.chain(freq) or []
        growth = any(s in ("yoy", "qoq", "qoq_ann") or s.startswith("diff") for s in chain)
        return "abs" if growth or self.rate else "pct"


@dataclass
class ReportPlan:
    """Everything scenario_report needs, checked against the data."""

    title: str
    baseline: str
    scenarios: List[str]
    groups: List[str]
    aggregates: Dict[str, Dict[str, float]] = field(default_factory=dict)
    quarterly: Optional[Tuple[str, str]] = None
    annual: Optional[Tuple[str, str]] = None
    variables: List[ReportVariable] = field(default_factory=list)
    decimals: int = 2
    meta: Optional[pd.DataFrame] = None
    history_end: Dict[Tuple[str, str], str] = field(default_factory=dict)
    mark_forecast: bool = False

    @property
    def forecasts(self) -> List[str]:
        """Baseline first, then the scenarios."""
        return [self.baseline] + self.scenarios

    @property
    def row_groups(self) -> List[str]:
        """Groups, then aggregates, in table order."""
        return self.groups + [a for a in self.aggregates if a not in self.groups]

    def period_range(self, freq: str) -> Optional[Tuple[str, str]]:
        return self.quarterly if freq == "Q" else self.annual


def _as_list(x) -> List[str]:
    if x is None:
        return []
    items = x.split(",") if isinstance(x, str) else list(x)
    return [str(i).strip() for i in items if str(i).strip()]


def _chain_or_none(x) -> Optional[List[str]]:
    if x is None or x is False or (isinstance(x, (list, tuple)) and not x):
        return None
    if isinstance(x, str) and x.strip().lower() in ("", "none", "no", "off"):
        return None
    return transforms.parse_steps(x)


def _is_annual_step(step: str) -> bool:
    return step == "annual" or step.startswith("annual_")


def _flag(x) -> bool:
    return x is True or str(x).strip().lower() in ("1", "true", "yes", "on")


def _deviation_setting(x, where: str) -> Optional[str]:
    if x is None or x is False:
        return None
    s = str(x).strip().lower()
    if s in ("", "none", "no", "off", "false"):
        return None
    if s in ("auto", "default"):
        return "auto"
    if s in ("abs", "absolute", "pp", "diff", "difference"):
        return "abs"
    if s in ("pct", "percent", "%", "relative"):
        return "pct"
    raise ValueError(f"{where}: deviation must be abs, pct or none; got {x!r}.")


def _table_prefix(raw: str, used: set) -> str:
    stem = _SHEET_BAD.sub("_", str(raw)).strip() or "VAR"
    stem = stem[:25]                                   # + "_A_dev" stays within Excel's 31 characters
    name, k = stem, 2
    while name.lower() in used:
        tag = f"_{k}"
        name, k = stem[: 25 - len(tag)] + tag, k + 1
    used.add(name.lower())
    return name


def parse_variables(spec: dict) -> List[ReportVariable]:
    """Read and check the `variables:` list of a report config.

    Args: spec is the report config dict. Returns ReportVariable objects with unique table
    name prefixes. Raises ValueError with a fix hint for missing indicators, unknown steps,
    deviation steps inside a chain or annual steps in a quarterly chain. A `transforms:`
    chain (older shorthand) fills `annual` when it has an annual step, else `quarterly`.
    """
    items = (spec or {}).get("variables")
    if not isinstance(items, list) or not items:
        raise ValueError("The report config needs `variables:`, a list such as "
                         "[{indicator: GDP, label: Real GDP growth, annual: [annual_avg, yoy], deviation: abs}].")
    out: List[ReportVariable] = []
    used: set = set()
    for i, item in enumerate(items):
        if not isinstance(item, dict) or not str(item.get("indicator") or "").strip():
            raise ValueError(f"variables[{i}] needs an indicator, e.g. {{indicator: GDP, annual: [annual_avg, yoy]}}.")
        ind = str(item["indicator"]).strip()
        where = f"variable {ind}"
        deviation = _deviation_setting(item.get("deviation", "auto"), where)
        q, a = _chain_or_none(item.get("quarterly")), _chain_or_none(item.get("annual"))
        if q is None and a is None and item.get("transforms") is not None:
            steps = transforms.parse_steps(item["transforms"])
            devs = [s for s in steps if s.startswith("dev_")]
            if devs and "deviation" not in item:
                deviation = devs[-1][4:]
            steps = [s for s in steps if not s.startswith("dev_")] or ["level"]
            if any(_is_annual_step(s) for s in steps):
                a = steps
            else:
                q = steps
        if q is None and a is None:
            raise ValueError(f"{where}: give a quarterly: and/or annual: chain, e.g. annual: [annual_avg, yoy].")
        for chain in (q, a):
            if chain and any(s.startswith("dev_") for s in chain):
                raise ValueError(f"{where}: put deviations in `deviation: abs|pct`, not in the chain {chain}.")
        if q and any(_is_annual_step(s) for s in q):
            raise ValueError(f"{where}: the quarterly chain {q} has an annual step; move it to `annual:`.")
        if a is not None and not any(_is_annual_step(s) for s in a):
            a = ["annual_avg"] + [s for s in a if s != "level"]
        units = str(item.get("units") or "").strip()
        decimals = item.get("decimals")
        out.append(ReportVariable(
            indicator=ind, label=str(item.get("label") or ind), name=_table_prefix(item.get("name") or ind, used),
            quarterly=q, annual=a, deviation=deviation, units=units,
            rate=units.lower() in _RATE_UNITS, decimals=None if decimals is None else int(decimals)))
    return out


def _period_bounds(value, freq: str) -> Optional[Tuple[str, str]]:
    if value is None or (isinstance(value, str) and value.strip().lower() in ("", "all")):
        return None
    if isinstance(value, bool) or (isinstance(value, int) and not 1000 <= value <= 9999):
        raise ValueError(f"Period range {value!r} is not readable: quote ranges in YAML, e.g. "
                         f"quarterly: \"2026Q1:2028Q4\" and annual: \"2026:2030\" (unquoted 2026:2030 is "
                         "read as a number).")
    start, end = parse_range(value)
    if freq_of(start) == freq:
        return start, end
    if freq == "Q":                                    # years given for the quarterly table
        return f"{parse_period(start)[0]}Q1", f"{parse_period(end)[0]}Q4"
    raise ValueError(f"periods.annual must be a range of years such as \"2026:2030\", got {value!r}.")


def plan_report(df: pd.DataFrame, spec: dict, *, meta: Optional[pd.DataFrame] = None) -> ReportPlan:
    """Check a report config against a long frame and resolve its defaults.

    Args: df is the long frame (baseline + scenarios); spec is the report config dict
    (`aggregates` may be inline or a path, resolved against spec["_base_dir"] or the
    current directory); meta is the variable metadata table (oegem.metadata), if any.
    Returns a ReportPlan. Raises ValueError when the baseline or a listed scenario is not
    in the data.
    """
    spec = spec or {}
    labels = [str(x) for x in pd.unique(df["forecast"])] if len(df) else []
    baseline = str(spec.get("baseline") or "baseline")
    if baseline not in labels:
        raise ValueError(f"Baseline {baseline!r} is not in the data (forecasts present: {', '.join(labels) or 'none'}); "
                         "set `baseline:` in the report config.")
    sel = spec.get("scenarios", "all")
    if sel is None or (isinstance(sel, str) and sel.strip().lower() == "all"):
        scenarios = [x for x in labels if x != baseline]
    else:
        scenarios = [x for x in _as_list(sel) if x != baseline]
        missing = [x for x in scenarios if x not in labels]
        if missing:
            raise ValueError(f"Scenario(s) not in the data: {', '.join(missing)}; present: {', '.join(labels)}.")
    agg_spec = spec.get("aggregates")
    aggregates = as_groupings(agg_spec, base_dir=spec.get("_base_dir")) if agg_spec else {}
    grp = spec.get("groups")
    if grp is None or (isinstance(grp, str) and grp.strip().lower() == "all"):
        groups = sorted(str(g) for g in pd.unique(df["group"]) if g not in aggregates)
    else:
        groups = _as_list(grp)
        absent = [g for g in groups if g not in set(df["group"])]
        if absent:
            log.warning("report: no data for group(s) %s", ", ".join(absent))
    periods = spec.get("periods") or {}
    if not isinstance(periods, dict):
        raise ValueError("`periods:` must be a mapping such as {quarterly: \"2026Q1:2028Q4\", annual: \"2026:2030\"}.")
    return ReportPlan(
        title=str(spec.get("title") or "Scenario report"), baseline=baseline, scenarios=scenarios,
        groups=groups, aggregates=aggregates,
        quarterly=_period_bounds(periods.get("quarterly"), "Q"), annual=_period_bounds(periods.get("annual"), "A"),
        variables=parse_variables(spec), decimals=int(spec.get("decimals", 2)),
        meta=meta, history_end=history_end(meta) if meta is not None else {},
        mark_forecast=_flag(spec.get("mark_forecast", False)))


# =========================================================================== tables
def _rank(order: Optional[List[str]], value) -> tuple:
    if order is not None and value in order:
        return (0, order.index(value), "")
    return (1, 0, str(value))


def table(df: pd.DataFrame, *, indicators=None, groups=None, forecasts=None, measure=None, start=None, end=None,
          rows: Sequence[str] = ("indicator", "group", "forecast"), decimals: Optional[int] = 2,
          freq: Optional[str] = None) -> pd.DataFrame:
    """Pivot a long frame into a table with one column per period.

    Args: df is a long frame; indicators, groups and forecasts filter rows and set their
    order; measure keeps one measure; start/end limit periods (inclusive); rows are the
    index levels; decimals rounds (None keeps full precision); freq ("Q" or "A") keeps
    one frequency. If several measures remain and "measure" is not a row level, it is
    added. Returns the table; columns sorted in time.
    """
    data = as_long_frame(df)
    order: Dict[str, List[str]] = {}
    for col, sel in (("indicator", indicators), ("group", groups), ("forecast", forecasts)):
        if sel is not None:
            order[col] = _as_list(sel)
            data = data[data[col].isin(order[col])]
    if measure is not None:
        data = data[data["measure"].isin(_as_list(measure))]
    if freq is not None:
        data = data[data["freq"] == str(freq).strip().upper()[:1]]
    if start is not None or end is not None:
        lo = normalize_period(start) if start is not None else None
        hi = normalize_period(end) if end is not None else None
        keep = {p: in_range(p, lo, hi) for p in pd.unique(data["period"])}
        data = data[data["period"].map(keep).astype(bool)]
    rows = [rows] if isinstance(rows, str) else list(rows)
    if "measure" not in rows and data["measure"].nunique() > 1 and data.duplicated(rows + ["period"]).any():
        rows.append("measure")
    if data.duplicated(rows + ["period"]).any():
        raise ValueError(f"Rows are not unique by {rows} and period; filter with measure=/freq= or add a row level.")
    if data.empty:
        return pd.DataFrame()
    wide = data.set_index(rows + ["period"])["value"].unstack("period")
    wide = wide[sorted(wide.columns, key=lambda p: (parse_period(p)[0], parse_period(p)[1] or 5))]
    keys = [k if isinstance(k, tuple) else (k,) for k in wide.index]
    keys.sort(key=lambda k: tuple(_rank(order.get(level), v) for level, v in zip(rows, k)))
    index = pd.MultiIndex.from_tuples(keys, names=rows) if len(rows) > 1 else pd.Index([k[0] for k in keys], name=rows[0])
    wide = wide.reindex(index)
    wide.columns.name = "period"
    return wide.round(int(decimals)) if decimals is not None else wide


def _aggregates_meaningful(measure: str, rate: bool) -> bool:
    steps = measure.split("|")
    return rate or any(s in ("yoy", "qoq", "qoq_ann", "dev_pct") for s in steps)


def _with_aggregates(frame: pd.DataFrame, plan: ReportPlan, var: ReportVariable) -> pd.DataFrame:
    if not plan.aggregates or frame.empty:
        return frame
    measure = str(frame["measure"].iloc[0])
    if not _aggregates_meaningful(measure, var.rate):
        log.warning("aggregating %s in measure %r: fine for rates and ratios, meaningless for levels in "
                    "different currencies (use a growth chain or deviation: pct)", var.indicator, measure)
    agg = weighted_mean(frame, plan.aggregates)
    return pd.concat([frame, agg], ignore_index=True) if len(agg) else frame


def _limit(frame: Optional[pd.DataFrame], plan: ReportPlan, freq: str) -> Optional[pd.DataFrame]:
    if frame is None or frame.empty:
        return frame
    keep = frame["group"].isin(plan.row_groups).to_numpy()
    bounds = plan.period_range(freq)
    if bounds:
        lo, hi = period_key(bounds[0]), period_key(bounds[1])
        keys = frame["period"].map({p: period_key(p) for p in pd.unique(frame["period"])}).to_numpy()
        keep &= (keys >= lo) & (keys <= hi)
    return frame[keep].reset_index(drop=True)


def variable_frames(df: pd.DataFrame, plan: ReportPlan, var: ReportVariable,
                    freq: str) -> Optional[Tuple[pd.DataFrame, Optional[pd.DataFrame]]]:
    """Long frames behind one variable's tables at one frequency.

    Args: df is the long frame of levels; plan comes from plan_report; var is one of
    plan.variables; freq is "Q" or "A". Returns (values, deviations) with aggregates
    appended and periods/groups limited to the plan; deviations is None when the variable
    has deviation: none. Returns None when the variable has no chain for this frequency.
    """
    chain = var.chain(freq)
    if chain is None:
        return None
    kind = var.deviation_kind(freq)
    sub = df[(df["indicator"] == var.indicator) & df["forecast"].isin(plan.forecasts)]
    if (sub["measure"] == "level").any():
        sub = sub[sub["measure"] == "level"]
    if sub.empty:
        return empty_frame(), (empty_frame() if kind else None)
    values = transforms.apply(sub, chain, meta=plan.meta)
    values = values[values["freq"] == freq].reset_index(drop=True)
    devs = transforms.deviation(values, plan.baseline, kind) if kind else None
    values = _with_aggregates(values, plan, var)
    if devs is not None:
        devs = _with_aggregates(devs, plan, var)
    return _limit(values, plan, freq), _limit(devs, plan, freq)


def _pivot(frame: pd.DataFrame, plan: ReportPlan, forecasts: List[str]) -> pd.DataFrame:
    present = set(frame["group"])
    groups = [g for g in plan.row_groups if g in present]
    wide = table(frame, groups=groups, forecasts=forecasts, rows=("group", "forecast"), decimals=None)
    full = pd.MultiIndex.from_product([groups, forecasts], names=["group", "forecast"])
    wide = wide.reindex(full)
    wide.columns.name = "period"
    return wide


def _mark_forecast(t: pd.DataFrame, plan: ReportPlan, indicator: str) -> None:
    """Append "f" to quarterly column headers later than the earliest history end of the table's groups."""
    groups = set(t.index.get_level_values("group")) if "group" in (t.index.names or []) else set()
    ends = [_end_key(p) for (ind, grp), p in plan.history_end.items() if ind == indicator and grp in groups]
    if not ends:
        return
    cut = min(ends)
    t.columns = pd.Index([f"{c}f" if period_key(c) > cut else c for c in t.columns], name=t.columns.name)
    t.attrs["forecast_marked"] = True


def _source_text(spec: dict) -> str:
    if spec.get("_source"):
        return str(spec["_source"])
    if spec.get("input"):
        return f"input {spec['input']}"
    if isinstance(spec.get("forecasts"), dict):
        return "forecasts: " + "; ".join(f"{k} = {v}" for k, v in spec["forecasts"].items())
    return "long frame passed to scenario_report"


def _readme(plan: ReportPlan, tables: Dict[str, pd.DataFrame], source: str, skipped: List[str]) -> pd.DataFrame:
    rows: List[Tuple[str, str, str]] = [
        ("Report", "Title", plan.title),
        ("Report", "Source", source),
        ("Report", "Baseline", plan.baseline),
        ("Report", "Scenarios", ", ".join(plan.scenarios) or "(none)"),
        ("Report", "Groups", ", ".join(plan.groups) or "(none)"),
    ]
    if plan.aggregates:
        text = "; ".join(f"{a} = " + ", ".join(f"{g} {w:g}" for g, w in m.items()) for a, m in plan.aggregates.items())
        rows.append(("Report", "Aggregates", text + " (weighted means over members with a value; weights renormalised)"))
    rows += [
        ("Report", "Quarterly periods", ":".join(plan.quarterly) if plan.quarterly else "all available"),
        ("Report", "Annual periods", ":".join(plan.annual) if plan.annual else "all available"),
        ("Report", "Created", _dt.date.today().isoformat()),
    ]
    steps: List[str] = []
    for name, t in tables.items():
        a = t.attrs
        rows.append(("Tables", name, f"{a['label']} ({a['indicator']}), {_FREQ_WORD[a['freq']]}: "
                                     f"{a['measure_label']} [measure {a['measure']}]"))
        for s in str(a["measure"]).split("|"):
            s = "diff" if s.startswith("diff") else s
            if s not in steps:
                steps.append(s)
    for name in skipped:
        rows.append(("Tables", name, "no data for this variable in the source; table left out"))
    if plan.meta is not None:
        notes = indicator_notes(plan.meta, groups=plan.groups or None)
        for ind in dict.fromkeys(v.indicator for v in plan.variables):
            rows.append(("Metadata", ind, notes.get(ind, "no metadata for this indicator")))
    for s in steps:
        if s in _NOTES:
            rows.append(("Notes", s, _NOTES[s]))
    if any(t.attrs.get("forecast_marked") for t in tables.values()):
        rows.append(("Notes", "f", _FORECAST_NOTE))
    rows += [
        ("Notes", "rows", "One row per (group, forecast); aggregates are appended after the groups."),
        ("Notes", "blank cells", "No value: missing data, an incomplete year, or a lag before the first period."),
        ("Notes", "mnemonics", "Indicator and group codes are those of the model catalogue; check them with "
                               "`oegem catalogue search`."),
    ]
    out = pd.DataFrame(rows, columns=["section", "item", "value"])
    out.attrs.update({"kind": "readme", "title": "About this report"})
    return out


def scenario_report(df: pd.DataFrame, spec: dict, *, source: Optional[str] = None,
                    meta: Optional[pd.DataFrame] = None) -> Dict[str, pd.DataFrame]:
    """Build the named tables of a report config from a long frame of levels.

    Args: df is a long frame (baseline + scenarios, measure "level"); spec is the report
    config dict (see the module docstring; input/forecasts/outputs are ignored here);
    source is the text shown as the source in the README (default: from spec); meta is
    the variable metadata table (oegem.metadata), used for the README metadata section,
    the `annual` transform step and `mark_forecast: true`.
    Returns {"README": ..., "GDP_Q": ..., "GDP_Q_dev": ..., ...}. Each data table has rows
    (group, forecast), one column per period, full precision values, and ``attrs`` with
    label, indicator, freq, kind ("values"/"deviation"), measure, measure_label, title and
    decimals for the writers.
    """
    data = as_long_frame(df)
    plan = plan_report(data, spec, meta=meta)
    tables: Dict[str, pd.DataFrame] = {}
    skipped: List[str] = []
    for var in plan.variables:
        for freq in ("Q", "A"):
            frames = variable_frames(data, plan, var, freq)
            if frames is None:
                continue
            for kind, frame, forecasts in (("values", frames[0], plan.forecasts),
                                           ("deviation", frames[1], plan.scenarios)):
                if frame is None or not forecasts:
                    continue
                name = f"{var.name}_{freq}" + ("_dev" if kind == "deviation" else "")
                if frame.empty:
                    log.warning("report: table %s is empty (no %s data for %s)", name, _FREQ_WORD[freq], var.indicator)
                    skipped.append(name)
                    continue
                t = _pivot(frame, plan, forecasts)
                measure = str(frame["measure"].iloc[0])
                words = transforms.measure_label(measure, rate=var.rate)
                t.attrs.update({
                    "kind": kind, "label": var.label, "indicator": var.indicator, "freq": freq,
                    "measure": measure, "measure_label": words, "units": var.units,
                    "title": f"{var.label}, {_FREQ_WORD[freq]}: {words}",
                    "decimals": plan.decimals if var.decimals is None else var.decimals})
                if freq == "Q" and plan.mark_forecast:
                    _mark_forecast(t, plan, var.indicator)
                tables[name] = t
    readme = _readme(plan, tables, source or _source_text(spec or {}), skipped)
    return {README: readme, **tables}


# ========================================================================== writers
def _as_tables(tables) -> Dict[str, pd.DataFrame]:
    if isinstance(tables, pd.DataFrame):
        return {"table": tables}
    if not isinstance(tables, dict):
        raise ValueError("Expected a dict of name -> DataFrame (as returned by scenario_report).")
    return tables


def _is_readme(t: pd.DataFrame) -> bool:
    return t.attrs.get("kind") == "readme"


def _flatten(t: pd.DataFrame) -> Tuple[pd.DataFrame, int]:
    """Table with its index turned into leading columns; returns (frame, number of index columns)."""
    if isinstance(t.index, pd.RangeIndex) and t.index.name is None:
        flat, n = t.copy(), 0
    else:
        n = t.index.nlevels
        flat = t.reset_index()
    flat.columns = [str(c) for c in flat.columns]
    return flat, n


def _decimals(t: pd.DataFrame, default: Optional[int]) -> Optional[int]:
    d = t.attrs.get("decimals", default)
    return None if d is None else int(d)


def _is_number(v) -> bool:
    return isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, bool)


def _missing(v) -> bool:
    if v is None or v is pd.NA:
        return True
    return isinstance(v, (float, np.floating)) and math.isnan(v)


def _fmt(v, decimals: Optional[int], thousands: bool = True) -> str:
    if _missing(v):
        return ""
    if not _is_number(v):
        return str(v)
    if decimals is None:
        return f"{float(v):,g}" if thousands else f"{float(v):g}"
    text = f"{float(v):{',' if thousands else ''}.{decimals}f}"
    return text[1:] if text.startswith("-") and not text.strip("-0.,") else text   # "-0.00" -> "0.00"


def _numeric_columns(flat: pd.DataFrame, n_index: int) -> List[bool]:
    return [i >= n_index and pd.api.types.is_numeric_dtype(flat[c]) for i, c in enumerate(flat.columns)]


def _sheet_title(name: str, used: set) -> str:
    stem = _SHEET_BAD.sub("_", str(name)).strip() or "Sheet"
    title, k = stem[:31], 2
    while title.lower() in used:
        tag = f"_{k}"
        title, k = stem[: 31 - len(tag)] + tag, k + 1
    used.add(title.lower())
    return title


def to_excel(tables, path, *, decimals: int = 2) -> Path:
    """Write tables to one workbook, one sheet per table (the README sheet first if present).

    Args: tables is a dict name -> DataFrame (or one DataFrame); path is the .xlsx file;
    decimals is the number format used when a table has no ``attrs["decimals"]``.
    Headers are bold, panes frozen below the header and right of the row labels, columns
    sized to their contents, NaN written as blank cells. Returns the path.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    tables = _as_tables(tables)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    bold = Font(bold=True)
    for name, t in tables.items():
        ws = wb.create_sheet(_sheet_title(name, used))
        flat, n_index = _flatten(t)
        dec = _decimals(t, decimals)
        numeric = _numeric_columns(flat, n_index)
        ws.append(list(flat.columns))
        for cell in ws[1]:
            cell.font = bold
        for row in flat.itertuples(index=False, name=None):
            ws.append([None if _missing(v) else (v.item() if isinstance(v, np.generic) else v) for v in row])
        fmt = "General" if dec is None else ("#,##0" if dec == 0 else "#,##0." + "0" * dec)
        for j, (col, is_num) in enumerate(zip(flat.columns, numeric), start=1):
            texts = [col] + [_fmt(v, dec) if is_num else ("" if _missing(v) else str(v)) for v in flat[col]]
            width = min(max(len(s) for s in texts) + 2, 60)
            if is_num:
                for (cell,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    cell.number_format = fmt
            if _is_readme(t) and col == "value":
                width = 100
                for (cell,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[get_column_letter(j)].width = max(width, 8)
        ws.freeze_panes = ws.cell(row=2, column=n_index + 1)
    if not tables:
        wb.create_sheet(README)
    wb.save(path)
    return path


def _file_stem(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", str(name)).strip("_") or "table"


def to_csv(tables, folder, *, decimals: Optional[int] = None) -> List[Path]:
    """Write one CSV per table into folder (``<name>.csv``, row labels as leading columns).

    Args: tables is a dict name -> DataFrame (or one DataFrame); folder is created if
    needed; decimals rounds tables without ``attrs["decimals"]`` (None: full precision).
    Returns the list of files written.
    """
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    files = []
    for name, t in _as_tables(tables).items():
        flat, _ = _flatten(t)
        dec = _decimals(t, decimals)
        path = folder / f"{_file_stem(name)}.csv"
        flat.to_csv(path, index=False, float_format=None if dec is None else f"%.{dec}f")
        files.append(path)
    return files


def _md_escape(text) -> str:
    return str(text).replace("\\", "\\\\").replace("|", "\\|").replace("\n", " ")


def _md_table(t: pd.DataFrame, decimals: Optional[int]) -> str:
    flat, n_index = _flatten(t)
    numeric = _numeric_columns(flat, n_index)
    lines = ["| " + " | ".join(_md_escape(c) for c in flat.columns) + " |",
             "|" + "|".join("---:" if num else ":---" for num in numeric) + "|"]
    for row in flat.itertuples(index=False, name=None):
        cells = [_fmt(v, decimals) if num else _md_escape("" if _missing(v) else v) for v, num in zip(row, numeric)]
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def to_markdown(table, path=None, decimals: Optional[int] = None, *, title: Optional[str] = None) -> str:
    """Render a table (or a dict of tables) as GitHub-flavoured markdown; no tabulate needed.

    Args: table is a DataFrame or a dict name -> DataFrame (a dict becomes a document with
    one section per table, the README first); path optionally receives the text (UTF-8);
    decimals overrides the tables' own setting (default 2); title heads a document.
    Returns the markdown text.
    """
    if isinstance(table, pd.DataFrame):
        text = _md_table(table, decimals if decimals is not None else _decimals(table, 2))
    else:
        parts = [f"# {title or 'Scenario report'}", ""]
        for name, t in _as_tables(table).items():
            if _is_readme(t):
                parts += [f"## {name}", ""]
                for section, block in t.groupby("section", sort=False):
                    body = block[["item", "value"]].reset_index(drop=True)
                    parts += [f"### {section}", "", _md_table(body, None), ""]
                continue
            parts += [f"## {name}", ""]
            if t.attrs.get("title"):
                parts += [_md_escape(t.attrs["title"]), ""]
            parts += [_md_table(t, decimals if decimals is not None else _decimals(t, 2)), ""]
        text = "\n".join(parts)
    if path is not None:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text + "\n", encoding="utf-8")
    return text


_CSS = """
:root { color-scheme: light; --page: #f9f9f7; --surface: #fcfcfb; --ink: #0b0b0b; --ink-2: #52514e;
  --muted: #6f6e69; --grid: #e1e0d9; --rule: #c3c2b7; --accent: #1c5cab; }
@media (prefers-color-scheme: dark) {
  :root:not([data-theme="light"]) { color-scheme: dark; --page: #0d0d0d; --surface: #1a1a19; --ink: #ffffff;
    --ink-2: #c3c2b7; --muted: #9c9a92; --grid: #2c2c2a; --rule: #4a4a45; --accent: #6da7ec; }
}
:root[data-theme="dark"] { color-scheme: dark; --page: #0d0d0d; --surface: #1a1a19; --ink: #ffffff;
  --ink-2: #c3c2b7; --muted: #9c9a92; --grid: #2c2c2a; --rule: #4a4a45; --accent: #6da7ec; }
* { box-sizing: border-box; }
body { margin: 0; background: var(--page); color: var(--ink);
  font: 14px/1.5 system-ui, -apple-system, "Segoe UI", sans-serif; }
main { max-width: 1180px; margin: 0 auto; padding: 32px 16px 64px; }
h1 { font-size: 24px; line-height: 1.25; margin: 0 0 4px; }
h2 { font-size: 17px; margin: 0 0 2px; }
h3 { font-size: 14px; margin: 16px 0 6px; color: var(--ink-2); }
.sub, .desc { color: var(--ink-2); margin: 0 0 12px; }
a { color: var(--accent); text-decoration: none; }
a:hover { text-decoration: underline; }
nav ul { list-style: none; padding: 0; margin: 16px 0 32px; display: flex; flex-wrap: wrap; gap: 4px 16px; }
section { margin: 0 0 36px; scroll-margin-top: 16px; }
.scroll { overflow-x: auto; background: var(--surface); border: 1px solid var(--grid); border-radius: 8px; }
table { border-collapse: collapse; width: 100%; font-variant-numeric: tabular-nums; }
th, td { padding: 6px 12px; border-bottom: 1px solid var(--grid); white-space: nowrap; }
thead th { text-align: right; color: var(--ink-2); font-weight: 600; border-bottom: 1px solid var(--rule); }
thead th.ix, tbody th { text-align: left; }
tbody th { font-weight: 500; }
td { text-align: right; }
tbody tr:last-child > * { border-bottom: 0; }
tr.first > * { border-top: 1px solid var(--rule); }
tbody tr:first-child > * { border-top: 0; }
.rep { color: var(--muted); font-weight: 400; }
table.readme td, table.readme th { text-align: left; white-space: normal; vertical-align: top; }
table.readme th { width: 14em; }
footer { color: var(--muted); font-size: 12px; margin-top: 40px; }
"""


def _slug(name: str, used: set) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", str(name).lower()).strip("-") or "table"
    slug, k = base, 2
    while slug in used:
        slug, k = f"{base}-{k}", k + 1
    used.add(slug)
    return slug


def _html_table(t: pd.DataFrame, decimals: Optional[int]) -> str:
    esc = _html.escape
    flat, n_index = _flatten(t)
    numeric = _numeric_columns(flat, n_index)
    left = ' class="ix"'
    head = "".join(f'<th scope="col"{left if i < n_index or not numeric[i] else ""}>{esc(c)}</th>'
                   for i, c in enumerate(flat.columns))
    body, prev = [], None
    for row in flat.itertuples(index=False, name=None):
        first_key = row[0] if n_index > 1 else None
        cls = ' class="first"' if n_index > 1 and first_key != prev else ""
        cells = []
        for i, v in enumerate(row):
            if i < n_index:
                rep = ' class="rep"' if i == 0 and n_index > 1 and first_key == prev else ""
                cells.append(f'<th scope="row"{rep}>{esc("" if _missing(v) else str(v))}</th>')
            elif numeric[i]:
                cells.append(f"<td>{esc(_fmt(v, decimals))}</td>")
            else:
                cells.append(f"<td>{esc('' if _missing(v) else str(v))}</td>")
        prev = first_key
        body.append(f"<tr{cls}>{''.join(cells)}</tr>")
    return f'<div class="scroll"><table><thead><tr>{head}</tr></thead><tbody>{"".join(body)}</tbody></table></div>'


def _html_readme(t: pd.DataFrame) -> str:
    esc = _html.escape
    parts = []
    for section, block in t.groupby("section", sort=False):
        rows = "".join(f'<tr><th scope="row">{esc(str(r.item))}</th><td>{esc(str(r.value))}</td></tr>'
                       for r in block.itertuples(index=False))
        parts.append(f'<h3>{esc(str(section))}</h3><div class="scroll"><table class="readme"><tbody>{rows}'
                     "</tbody></table></div>")
    return "".join(parts)


def to_html(tables, path, title: str = "Scenario report", *, decimals: Optional[int] = None) -> Path:
    """Write tables to one self-contained HTML page (inline CSS, light and dark themes).

    Args: tables is a dict name -> DataFrame (or one DataFrame); path is the .html file;
    title heads the page; decimals overrides the tables' own setting (default 2).
    One section per table with its description; a contents list links them. Returns the path.
    """
    esc = _html.escape
    tables = _as_tables(tables)
    used: set = set()
    nav, sections = [], []
    for name, t in tables.items():
        anchor = _slug(name, used)
        nav.append(f'<li><a href="#{anchor}">{esc(str(name))}</a></li>')
        if _is_readme(t):
            inner = _html_readme(t)
        else:
            desc = f'<p class="desc">{esc(str(t.attrs["title"]))}</p>' if t.attrs.get("title") else ""
            inner = desc + _html_table(t, decimals if decimals is not None else _decimals(t, 2))
        sections.append(f'<section id="{anchor}"><h2>{esc(str(name))}</h2>{inner}</section>')
    page = (
        "<!doctype html>\n<html lang=\"en\">\n<head>\n<meta charset=\"utf-8\">\n"
        "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">\n"
        f"<title>{esc(title)}</title>\n<style>{_CSS}</style>\n</head>\n<body>\n<main>\n"
        f"<h1>{esc(title)}</h1>\n<p class=\"sub\">{len([t for t in tables.values() if not _is_readme(t)])} tables"
        f" &middot; {_dt.date.today().isoformat()}</p>\n"
        f"<nav aria-label=\"Tables\"><ul>{''.join(nav)}</ul></nav>\n" + "\n".join(sections) +
        "\n<footer>Blank cells have no value. Check indicator and group codes in your own catalogue "
        "(<code>oegem catalogue search</code>).</footer>\n</main>\n</body>\n</html>\n")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(page, encoding="utf-8")
    return path


# ===================================================================== build_report
def load_report_config(config, base_dir=None) -> Tuple[dict, Path]:
    """Read a report config (YAML/JSON path or dict) and return (spec, base directory).

    Relative paths in the config resolve against the config file's folder (for a dict:
    base_dir, default the current directory). The base directory is also stored in
    spec["_base_dir"]. Raises ValueError for a missing file or a malformed config.
    """
    if isinstance(config, (str, Path)):
        path = Path(config)
        if not path.is_file():
            raise ValueError(f"Report config not found: {path}")
        spec = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
        base = path.resolve().parent
    elif isinstance(config, dict):
        spec = copy.deepcopy(config)
        base = Path(base_dir or spec.get("_base_dir") or Path.cwd())
    else:
        raise ValueError("Pass a report config file path or a dict.")
    if not isinstance(spec, dict):
        raise ValueError("A report config must be a mapping with input: (or forecasts:) and variables:.")
    has_input, has_fc = bool(spec.get("input")), bool(spec.get("forecasts"))
    if has_input == has_fc:
        raise ValueError("Give exactly one of `input:` (a run directory or results file) and `forecasts:` "
                         "(label -> release spec or workspace path) in the report config.")
    parse_variables(spec)                                    # fail early on a bad variables list
    spec["_base_dir"] = str(base)
    return spec, base


def _resolve(base: Path, p) -> Path:
    path = Path(str(p))
    # resolve() folds ".." segments so deep folders stay under Windows' path limit
    return (path if path.is_absolute() else base / path).resolve()


def _load_input(path: Path) -> Tuple[pd.DataFrame, dict, Path]:
    from .data import load

    manifest: dict = {}
    if path.is_dir():
        found = [path / n for n in RESULT_FILES if (path / n).is_file()]
        if not found:
            raise ValueError(f"No results file ({', '.join(RESULT_FILES)}) in {path}; run `oegem run CONFIG` "
                             "first or point `input:` at a results file.")
        results, mpath = found[0], path / "manifest.json"
    elif path.is_file():
        results, mpath = path, path.parent / "manifest.json"
    else:
        raise ValueError(f"Report input not found: {path}")
    if mpath.is_file():
        try:
            manifest = json.loads(mpath.read_text(encoding="utf-8-sig"))
        except ValueError:
            log.warning("could not read %s; continuing without it", mpath)
    return load(results), (manifest if isinstance(manifest, dict) else {}), results


def load_results(path) -> pd.DataFrame:
    """Long frame of a run directory (its results.csv/.parquet/.xlsx) or of a results file."""
    return _load_input(Path(path))[0]


def _manifest_note(m: dict) -> str:
    bits = []
    if m.get("run_name"):
        bits.append(f"run {m['run_name']}")
    rel = m.get("release")
    name = rel.get("name") if isinstance(rel, dict) else (rel if isinstance(rel, str) else None)
    name = name or m.get("release_name")
    if name:
        bits.append(f"release {name}")
    return "; ".join(bits)


def _fetch_forecasts(spec: dict, client, *,
                     with_meta: bool = True) -> Tuple[pd.DataFrame, str, Optional[pd.DataFrame]]:
    from .data import fetch_many
    from .resources import GEM_ROOT, resolve_release

    forecasts = spec["forecasts"]
    if not isinstance(forecasts, dict):
        raise ValueError("`forecasts:` must map labels to release specs or workspace paths, e.g. "
                         "{baseline: latest:5yr, my_run: /me/oegem-runs/x/y-20261015-101500}.")
    fetch_cfg = spec.get("fetch") or {}
    if client is None:
        from .client import OEClient
        client = OEClient.from_env()
    root = fetch_cfg.get("root") or (spec.get("model") or {}).get("root") or GEM_ROOT
    indicators = fetch_cfg.get("indicators") or list(dict.fromkeys(v.indicator for v in parse_variables(spec)))
    groups = fetch_cfg.get("groups")
    if not groups:
        cfg_groups = spec.get("groups")
        if isinstance(cfg_groups, str) and cfg_groups.strip().lower() == "all":
            groups = "all"
        else:
            groups = _as_list(cfg_groups)
            agg = spec.get("aggregates")
            if agg:
                for members in as_groupings(agg, base_dir=spec.get("_base_dir")).values():
                    groups += [g for g in members if g not in groups]
        if not groups:
            raise ValueError("With `forecasts:` give the groups to fetch: `fetch: {groups: [...]}` or `groups:`.")
    ids, notes = {}, []
    for label, rel in forecasts.items():
        res = resolve_release(client, rel, root=root)
        fid = res.get("Id") or res.get("id")
        if not fid:
            raise ValueError(f"No forecast id found for {label}: {rel!r}; check the path with `oegem releases`.")
        ids[str(label)] = fid
        shown = res.get("Name")
        notes.append(f"{label} = {rel}" + (f" ({shown})" if shown and str(shown) != str(rel) else ""))
    df = fetch_many(client, ids, indicators, groups, start=fetch_cfg.get("from"), end=fetch_cfg.get("to"),
                    freq=fetch_cfg.get("freq"), annual_fallback=bool(fetch_cfg.get("annual_fallback", True)))
    meta = None
    if with_meta and len(df):
        fid = ids.get(str(spec.get("baseline") or "baseline")) or next(iter(ids.values()))
        pairs = list(dict.fromkeys(zip(df["indicator"].astype(str), df["group"].astype(str))))
        try:
            meta = fetch_metadata(client, fid, pairs=pairs)
        except Exception as exc:  # metadata is an extra: never fail the report over it
            log.warning("variable metadata not fetched (%s); the report goes on without it", exc)
    return df, "forecasts: " + "; ".join(notes), meta


def _run_metadata(folder: Path) -> Optional[pd.DataFrame]:
    """The metadata.csv of a run folder, or None (absent or unreadable)."""
    path = Path(folder) / METADATA_FILE
    if not path.is_file():
        return None
    try:
        return load_metadata(path)
    except (OSError, ValueError) as exc:
        log.warning("could not read %s (%s); the report goes on without metadata", path, exc)
        return None


def build_report(report_config_path_or_dict, *, client=None, base_dir=None) -> dict:
    """Run a report config end to end: load or fetch data, build tables, write the outputs.

    Args: report_config_path_or_dict is a YAML/JSON path or a dict (see the module
    docstring); client is used for `forecasts:` (default OEClient.from_env(), which is the
    demo backend when OEGEM_DEMO=1); base_dir resolves relative paths of a dict config
    (default: the current directory).
    Returns {"tables": {name: DataFrame}, "files": [Path, ...]}. Outputs go to
    outputs.dir as report.xlsx, csv/<table>.csv, report.md, report.html and
    report_charts.pdf, as switched on under `outputs:` (default: xlsx only). Variable
    metadata comes from the run folder's metadata.csv or, with `forecasts:`, from the
    baseline (`metadata: false` skips it); see the module docstring.
    """
    spec, base = load_report_config(report_config_path_or_dict, base_dir)
    use_meta = spec.get("metadata", True) is not False
    meta = None
    if spec.get("input"):
        df, manifest, results = _load_input(_resolve(base, spec["input"]))
        kind = "run directory" if results.parent == _resolve(base, spec["input"]) else "results file"
        note = _manifest_note(manifest)
        source = f"{kind} {spec['input']}" + (f" ({note})" if note else "")
        if use_meta:
            meta = _run_metadata(results.parent)
    else:
        df, source, meta = _fetch_forecasts(spec, client, with_meta=use_meta)
    if spec.get("aggregates"):
        spec["aggregates"] = as_groupings(spec["aggregates"], base_dir=base)
    tables = scenario_report(df, spec, source=source, meta=meta)
    out_cfg = spec.get("outputs") or {}
    if not isinstance(out_cfg, dict):
        raise ValueError("`outputs:` must be a mapping such as {dir: report, xlsx: true, html: true}.")
    out_dir = _resolve(base, out_cfg.get("dir") or "report")
    stem = _file_stem(out_cfg.get("name") or "report")
    title = str(spec.get("title") or "Scenario report")
    files: List[Path] = []
    if out_cfg.get("xlsx", True):
        files.append(to_excel(tables, out_dir / f"{stem}.xlsx"))
    if out_cfg.get("csv", False):
        files.extend(to_csv(tables, out_dir / "csv"))
    if out_cfg.get("markdown", False):
        to_markdown(tables, out_dir / f"{stem}.md", title=title)
        files.append(out_dir / f"{stem}.md")
    if out_cfg.get("html", False):
        files.append(to_html(tables, out_dir / f"{stem}.html", title))
    if out_cfg.get("charts", False):
        from .plots import chart_book
        files.append(chart_book(df, spec, out_dir / f"{stem}_charts.pdf", meta=meta))
    log.info("report: %d tables, %d files in %s", len(tables) - 1, len(files), out_dir)
    return {"tables": tables, "files": files}
